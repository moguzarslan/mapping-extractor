#!/usr/bin/env python3
"""
Causes of isPartOf errors in an architecture evaluation, pooled over every
document and run of one version.

The architecture evaluator scores isPartOf as an Accuracy over the matched pairs in
which the LLM populated the field (service/architecture_evaluator_service.py):
  - units / patterns: the LLM parent ids, mapped to GT elements THROUGH the element
    matching, must equal the GT parent set exactly;
  - connectors: the hub-centred endpoint pairs, mapped through the matching, must
    equal the GT pairs.

For every document and run the script reads the matched pairs the evaluator
produced (Matched_TP of <doc>_arch_eval.xlsx), the LLM extraction JSON and the
ground-truth workbook, re-judges every pair with the evaluator's own
`field_agrees` (no embeddings are computed, so the figures reproduce the reported
ones) and attributes every wrong pair to ONE primary cause, checked in this order:

  Ground truth   GT has no parent / GT parent id does not resolve
  Upstream       a GT parent (connector endpoint) was never extracted or matched
  Reference      the LLM referenced an id that does not exist in its own output
  Upstream       the LLM linked to one of its own false-positive elements
  Evaluator      connector with the right endpoints but a different first (hub) one
  Linking        every element involved was matched, but the LLM chose the wrong
                 parent(s): missing, extra, wrong class, wrong type, wrong element

Output workbook (analysis/report/isPartOf_causes_<version>.xlsx), two sheets,
every document and run pooled:
  Causes           every wrong pair by cause, with its share of all errors
  Errors_By_Class  errors by class of the scored element (unit / pattern /
                   connector) and their share of errors, EXCLUDING the pairs whose
                   GT parent or GT connector endpoint was never extracted — the
                   errors no isPartOf instruction can fix

Usage (from any directory):
    analysis/script/ispartof_causes.py                      # design/v3
    analysis/script/ispartof_causes.py --version design/v2
    analysis/script/ispartof_causes.py --output some/other.xlsx
"""
import os
import subprocess
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "analysis" / "report"
VENV_PYTHON = PROJECT_ROOT / ".venv" / "bin" / "python"
_REEXEC_FLAG = "ISPARTOF_CAUSES_REEXEC"


def _sysctl(name: str) -> str:
    try:
        return subprocess.run(["/usr/sbin/sysctl", "-n", name],
                              capture_output=True, text=True).stdout.strip()
    except OSError:
        return ""


def _ensure_project_interpreter() -> None:
    """Re-run the script under the project's virtualenv, natively.

    The evaluator needs the virtualenv (Python 3.10+ syntax, its packages), so a
    script started by another interpreter — e.g. through the shebang — is re-run
    under it. On Apple Silicon that goes through `arch -arm64`: a process running
    under Rosetta (such as an Intel Anaconda python3) otherwise hands x86_64 on to
    the interpreter it execs, and the virtualenv's arm64 packages fail to load.
    """
    if os.environ.get(_REEXEC_FLAG) or not VENV_PYTHON.exists():
        return
    in_venv = Path(sys.prefix).resolve() == VENV_PYTHON.parent.parent.resolve()
    translated = sys.platform == "darwin" and _sysctl("sysctl.proc_translated") == "1"
    if in_venv and not translated:
        return
    command = [str(VENV_PYTHON), str(Path(__file__).resolve()), *sys.argv[1:]]
    if sys.platform == "darwin" and _sysctl("hw.optional.arm64") == "1":
        command = ["/usr/bin/arch", "-arm64", *command]
    os.environ[_REEXEC_FLAG] = "1"
    os.execv(command[0], command)


_ensure_project_interpreter()

import argparse  # noqa: E402

import pandas as pd  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

sys.path.insert(0, str(PROJECT_ROOT))
from service import architecture_evaluator_service as ev  # noqa: E402

THRESHOLD = 0.75  # unused by the isPartOf judgement, required by the signature

GROUND_TRUTH = "Ground truth"
UPSTREAM = "Upstream (extraction / matching)"
REFERENCE = "Reference error"
EVALUATOR = "Evaluator rule"
LINKING = "Linking (wrong parent chosen)"

GT_PARENT_NOT_EXTRACTED = "GT parent not extracted / not matched"
GT_ENDPOINT_NOT_EXTRACTED = "GT endpoint not extracted / not matched"
# Errors no isPartOf instruction can fix: the correct parent does not exist in the
# LLM output, so the pair is wrong whatever the model links.
NOT_EXTRACTED_CAUSES = {GT_PARENT_NOT_EXTRACTED, GT_ENDPOINT_NOT_EXTRACTED}

CLASSES = ["unit", "pattern", "connector"]


# ---------------------------------------------------------------------------
# Cause attribution
# ---------------------------------------------------------------------------
def resolve_refs(rec: dict, index: dict) -> tuple[list[int], list[str]]:
    """An element's isPartOf ids resolved within its own side: (indices in order,
    without duplicates; ids that do not resolve)."""
    resolved, dangling = [], []
    for ref in ev.to_id_list(rec.get("isPartOf")):
        idx = index.get(str(ref).strip())
        if idx is None:
            dangling.append(str(ref))
        elif idx not in resolved:
            resolved.append(idx)
    return resolved, dangling


def parent_kind(rec: dict) -> str:
    """`pattern`, `connector` or `unit:<Type>` — the level a parent sits at."""
    cls = ev.match_class(rec)
    return f"unit:{rec.get('type')}" if cls == "unit" else cls


def named_cause(l_rec, g_rec, gt, llm_idx, gt_idx, llm_to_gt, gt_to_llm):
    gt_par, gt_dangling = resolve_refs(g_rec, gt_idx)
    llm_par, llm_dangling = resolve_refs(l_rec, llm_idx)
    if not gt_par:
        if gt_dangling:
            return GROUND_TRUTH, "GT parent id does not resolve"
        return GROUND_TRUTH, "GT has no parent (LLM added one)"
    if any(g not in gt_to_llm for g in gt_par):
        return UPSTREAM, GT_PARENT_NOT_EXTRACTED
    if llm_dangling:
        return REFERENCE, "LLM parent id does not resolve"
    if any(i not in llm_to_gt for i in llm_par):
        return UPSTREAM, "LLM parent is its own false-positive element"
    if not llm_par:
        return LINKING, "LLM gave no usable parent"

    pred, true = {llm_to_gt[i] for i in llm_par}, set(gt_par)
    if pred < true:
        return LINKING, "missing parent(s)"
    if pred > true:
        return LINKING, "extra parent(s)"
    if {ev.match_class(gt[g]) for g in pred} != {ev.match_class(gt[g]) for g in true}:
        return LINKING, "wrong parent class (pattern / unit / connector)"
    if {parent_kind(gt[g]) for g in pred} != {parent_kind(gt[g]) for g in true}:
        return LINKING, "wrong parent type (same class)"
    return LINKING, "wrong parent element (same type)"


def connector_cause(l_rec, g_rec, gt, llm_idx, gt_idx, llm_to_gt, gt_to_llm):
    gt_ep, _ = resolve_refs(g_rec, gt_idx)
    llm_ep, llm_dangling = resolve_refs(l_rec, llm_idx)
    if len(gt_ep) < 2:
        return GROUND_TRUTH, "GT connector has < 2 resolvable endpoints"
    if any(g not in gt_to_llm for g in gt_ep):
        return UPSTREAM, GT_ENDPOINT_NOT_EXTRACTED
    if llm_dangling:
        return REFERENCE, "LLM endpoint id does not resolve"
    if any(i not in llm_to_gt for i in llm_ep):
        return UPSTREAM, "LLM endpoint is its own false-positive element"
    if len(llm_ep) < 2:
        return LINKING, "LLM connector has < 2 endpoints"

    pred, true = {llm_to_gt[i] for i in llm_ep}, set(gt_ep)
    if pred == true:
        return EVALUATOR, "same endpoints, different hub (first endpoint)"
    if pred < true:
        return LINKING, "missing endpoint(s)"
    if pred > true:
        return LINKING, "extra endpoint(s)"
    return LINKING, "wrong endpoint(s)"


# ---------------------------------------------------------------------------
# One run
# ---------------------------------------------------------------------------
def analyse_run(doc: str, run: str, gt_path: Path, llm_path: Path, eval_path: Path) -> list[dict]:
    ispartof = ev.SPEC_BY_NAME["isPartOf"]
    gt = ev.load_ground_truth(str(gt_path))
    llm = ev.load_llm_extraction(str(llm_path))
    llm_idx, gt_idx = ev.build_id_index(llm), ev.build_id_index(gt)

    matched = pd.read_excel(eval_path, sheet_name="Matched_TP", dtype=str)
    pairs = []
    for llm_id, gt_id in zip(matched.get("LLM_ID", []), matched.get("GT_ID", [])):
        i, j = llm_idx.get(str(llm_id).strip()), gt_idx.get(str(gt_id).strip())
        if i is None or j is None:
            print(f"  ! {doc}/{run}: matched pair {llm_id} / {gt_id} not found in inputs")
            continue
        pairs.append((i, j))
    llm_to_gt = {i: j for i, j in pairs}
    gt_to_llm = {j: i for i, j in pairs}

    rows = []
    for i, j in pairs:
        l_rec, g_rec = llm[i], gt[j]
        child_class = ev.match_class(g_rec)
        llm_pop = ev.field_present(l_rec, ispartof)
        gt_pop = ev.field_present(g_rec, ispartof)
        agrees = ev.field_agrees(ispartof, l_rec, g_rec, THRESHOLD, llm_idx=llm_idx,
                                 gt_idx=gt_idx, llm_to_gt=llm_to_gt)[0]
        correct = bool(llm_pop and gt_pop and agrees)

        group = cause = None
        if llm_pop and not correct:
            attribute = connector_cause if child_class == "connector" else named_cause
            group, cause = attribute(l_rec, g_rec, gt, llm_idx, gt_idx, llm_to_gt, gt_to_llm)
        rows.append({"child_class": child_class, "scored": llm_pop, "correct": correct,
                     "cause_group": group, "cause": cause})

    # Check against the evaluator's own counts for this run.
    counts = pd.read_excel(eval_path, sheet_name="Field_Counts")
    reported = counts[counts["field"] == "isPartOf"].iloc[0]
    scored = sum(r["scored"] for r in rows)
    correct = sum(r["correct"] for r in rows)
    if (scored, correct) != (int(reported["llm_populated_matched"]), int(reported["correct_in_matched"])):
        print(f"  WARNING {doc}/{run}: recomputed {correct}/{scored} differs from reported "
              f"{reported['correct_in_matched']}/{reported['llm_populated_matched']}")
    return rows


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------
def causes_sheet(pairs: pd.DataFrame) -> pd.DataFrame:
    wrong = pairs[pairs["scored"] & ~pairs["correct"]]
    table = (wrong.groupby(["cause_group", "cause"]).size()
             .reset_index(name="wrong_pairs")
             .sort_values("wrong_pairs", ascending=False))
    table["share_of_errors"] = (table["wrong_pairs"] / len(wrong)).round(4)
    table["counted_in_Errors_By_Class"] = ~table["cause"].isin(NOT_EXTRACTED_CAUSES)
    total = {"cause_group": "Total", "cause": "", "wrong_pairs": len(wrong),
             "share_of_errors": 1.0 if len(wrong) else 0.0, "counted_in_Errors_By_Class": ""}
    return pd.concat([table, pd.DataFrame([total])], ignore_index=True)


def errors_by_class_sheet(pairs: pd.DataFrame) -> pd.DataFrame:
    scored = pairs[pairs["scored"]]
    excluded = scored["cause"].isin(NOT_EXTRACTED_CAUSES)
    considered = scored[~excluded]
    total_wrong = int((~considered["correct"]).sum())

    def row(label, mask_scored, mask_considered):
        sub = considered[mask_considered]
        wrong = int((~sub["correct"]).sum())
        return {
            "class": label,
            "scored_pairs": int(mask_scored.sum()),
            "excluded (GT parent/endpoint not extracted)": int((excluded & mask_scored).sum()),
            "pairs_considered": len(sub),
            "correct": int(sub["correct"].sum()),
            "wrong": wrong,
            "accuracy": round(sub["correct"].sum() / len(sub), 4) if len(sub) else None,
            "share_of_errors": round(wrong / total_wrong, 4) if total_wrong else 0.0,
        }

    rows = [row(c, scored["child_class"] == c, considered["child_class"] == c) for c in CLASSES]
    rows.append(row("Total", scored["child_class"].notna(), considered["child_class"].notna()))
    return pd.DataFrame(rows)


def write_workbook(path: Path, sheets: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for cell in ws[ws.max_row]:  # Total row
                cell.font = Font(bold=True)
            for column in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
                ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--version", default="design/v3",
                        help="subdirectory under outputs/{evaluation,gemini}/architecture (default: design/v3)")
    parser.add_argument("--output", type=Path, default=None,
                        help="output xlsx (default: analysis/report/isPartOf_causes_<version>.xlsx)")
    args = parser.parse_args()

    eval_root = PROJECT_ROOT / "outputs" / "evaluation" / "architecture" / args.version
    llm_root = PROJECT_ROOT / "outputs" / "gemini" / "architecture" / args.version
    gt_root = PROJECT_ROOT / "resource" / "groundTruths" / "architecture"
    output = args.output or REPORT_DIR / f"isPartOf_causes_{args.version.replace('/', '_')}.xlsx"

    rows = []
    for eval_path in sorted(eval_root.glob("*/run_*/*_arch_eval.xlsx")):
        run, doc = eval_path.parent.name, eval_path.parent.parent.name
        llm_path = llm_root / doc / run / f"{doc}_architecture.json"
        gt_path = gt_root / f"{doc}_ground_truth_architecture.xlsx"
        if not llm_path.exists() or not gt_path.exists():
            print(f"Skipping {doc}/{run}: missing {llm_path if not llm_path.exists() else gt_path}")
            continue
        print(f"Analysing {doc}/{run}")
        rows += analyse_run(doc, run, gt_path, llm_path, eval_path)

    if not rows:
        raise SystemExit(f"No evaluated runs found under {eval_root}")

    pairs = pd.DataFrame(rows)
    write_workbook(output, {
        "Causes": causes_sheet(pairs),
        "Errors_By_Class": errors_by_class_sheet(pairs),
    })
    print(f"isPartOf causes saved: {output}")


if __name__ == "__main__":
    main()
